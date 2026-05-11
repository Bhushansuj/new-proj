import argparse
import sys
from pathlib import Path

import mysql.connector
from dotenv import load_dotenv
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = PROJECT_ROOT / "backend"


def env(name):
    import os

    return os.getenv(name)


def get_connection():
    load_dotenv(BACKEND_DIR / ".env")
    return mysql.connector.connect(
        host=env("DB_HOST"),
        user=env("DB_USER"),
        password=env("DB_PASSWORD"),
        database=env("DB_NAME"),
    )


def ensure_columns(cursor):
    cursor.execute("ALTER TABLE systems ADD COLUMN IF NOT EXISTS os_version VARCHAR(50) NULL")
    cursor.execute("ALTER TABLE systems ADD COLUMN IF NOT EXISTS gpu VARCHAR(100) NULL")
    cursor.execute("ALTER TABLE systems ADD COLUMN IF NOT EXISTS current_usage VARCHAR(255) NULL")
    cursor.execute("ALTER TABLE systems ADD COLUMN IF NOT EXISTS reserved_until VARCHAR(100) NULL")
    cursor.execute("ALTER TABLE systems ADD COLUMN IF NOT EXISTS reinstalled_at VARCHAR(100) NULL")
    cursor.execute("ALTER TABLE systems ADD COLUMN IF NOT EXISTS info TEXT NULL")
    cursor.execute("ALTER TABLE systems ADD COLUMN IF NOT EXISTS helpline TEXT NULL")
    cursor.execute("ALTER TABLE system_status ADD UNIQUE KEY IF NOT EXISTS uq_system_status_system_id (system_id)")


def clean(value):
    if value is None:
        return None

    value = str(value).strip()
    return value or None


def usage_status(current_usage, reserved_until):
    usage = (current_usage or "").strip().lower()

    if reserved_until:
        return "reserved"

    if not usage or usage == "frei":
        return "free"

    return "busy"


def import_workbook(path):
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    connection = get_connection()
    cursor = connection.cursor()
    ensure_columns(cursor)

    current_type = None
    imported = 0

    for row in sheet.iter_rows(min_row=4, values_only=True):
        hostname = clean(row[0])

        if not hostname:
            continue

        if hostname == "Test-Rechner":
            current_type = "PC"
            continue

        if hostname == "Test-Laptops":
            current_type = "Laptop"
            continue

        if hostname == "Test-VMs":
            current_type = "VM"
            continue

        mac_address = clean(row[1])
        os_version = clean(row[2])
        model = clean(row[3])
        gpu = clean(row[4])
        location = clean(row[5])
        current_usage = clean(row[6])
        reserved_until = clean(row[7])
        reinstalled_at = clean(row[8])
        info = clean(row[9])
        helpline = clean(row[10]) if len(row) > 10 else None

        cursor.execute(
            """
            INSERT INTO systems
            (
                hostname, system_type, mac_address, model, location, inventory_number,
                os_version, gpu, current_usage, reserved_until, reinstalled_at, info, helpline
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE
                system_type = VALUES(system_type),
                mac_address = VALUES(mac_address),
                model = VALUES(model),
                location = VALUES(location),
                os_version = VALUES(os_version),
                gpu = VALUES(gpu),
                current_usage = VALUES(current_usage),
                reserved_until = VALUES(reserved_until),
                reinstalled_at = VALUES(reinstalled_at),
                info = VALUES(info),
                helpline = VALUES(helpline)
            """,
            (
                hostname,
                current_type or "PC",
                mac_address,
                model,
                location,
                None,
                os_version,
                gpu,
                current_usage,
                reserved_until,
                reinstalled_at,
                info,
                helpline,
            ),
        )
        cursor.execute("SELECT id FROM systems WHERE hostname = %s", (hostname,))
        system_id = cursor.fetchone()[0]
        cursor.execute(
            """
            INSERT INTO system_status (system_id, power_status, usage_status)
            VALUES (%s, 'offline', %s)
            ON DUPLICATE KEY UPDATE usage_status = VALUES(usage_status)
            """,
            (system_id, usage_status(current_usage, reserved_until)),
        )
        imported += 1

    connection.commit()
    cursor.close()
    connection.close()
    return imported


def main():
    parser = argparse.ArgumentParser(description="Import NDR Testcenter Excel systems into MySQL.")
    parser.add_argument("xlsx", help="Path to the NDR Testcenter Excel file")
    args = parser.parse_args()

    count = import_workbook(args.xlsx)
    print(f"{count} systems imported or updated.")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"Import failed: {error}", file=sys.stderr)
        raise
